


from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

sections = [
    ("Introduction to GIS", [
        "What GIS is and why it’s important",
        "How our organization uses GIS",
        "OpenGov’s role as our GIS partner"
    ]),
    ("GIS Basics", []),
    ("Choosing and Accessing Maps", [
        "Which map to use for different tasks",
        "How to log in"
    ]),
    ("Basemaps", ["Purpose of basemaps", "When to switch basemaps"]),
    ("Layers", [
        "Common layers employees will use",
        "Turning layers on and off",
        "Understanding symbols and color coding",
        "Meaning of specific layers: Projects, Connects, Omni/Trace Wire boxes, Service Inquiries, Warnings, Open Data, Survey Points"
    ]),
    ("Searching and Navigation", [
        "Search by address/road",
        "Search by meter number",
        "Search by premise number",
        "Search by SEMS ID",
        "Navigation tools"
    ]),
    ("Understanding Locations", [
        "What is an RD?",
        "What is an RRD?",
        "Meaning of Customer Service locations"
    ]),
    ("Repair Markups", [
        "How to complete markups",
        "Examples of proper drawings",
        "Importance of clear notes"
    ]),
    ("Photos and Documentation", [
        "Examples of helpful photos",
        "Examples of unhelpful photos"
    ]),
    ("Q&A", ["Open discussion and questions"])
]

for title, bullets in sections:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    for b in bullets:
        p = body.add_paragraph()
        p.text = b
        p.level = 0

prs.save("New_Employee_GIS_Training.pptx")
"File created"
